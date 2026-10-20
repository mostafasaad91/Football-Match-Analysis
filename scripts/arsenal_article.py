"""Write the Arsenal article as a Word file next to the graphics it walks through.

The text is fixed prose about one block of rounds (Premier League 2026-27, rounds 1-5); the
figures in it were read off the aggregate tables, and the images come from
``team_story.build_everything``: every club graphic of the block with Arsenal picked out.
Run after the images exist:

    python scripts/arsenal_article.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

FOLDER = Path("output/aggregates/England_Premier_League_2026-2027_R1-5/article_arsenal_full")
TITLE = "Arsenal Defend With the Ball. Five Games Show What Happens When They Can't."
STANDFIRST = (
    "Twelve points from fifteen, the second-meanest defence in the league by chances allowed, "
    "and two away games that gave up three times as much as the other three put together. "
    "Every chart from the first five rounds, read for what it says about one team."
)

# A heading starts a chapter. Every other entry is (image prefix, caption, text): the text is
# printed first and leads up to the image, so each graphic is introduced by the one before it.
H = "heading"
BLOCKS = [
    (
        "club_profile",
        "Arsenal's profile after five rounds. Slice length is the share of the league they beat on that measure.",
        "Second place after five rounds is the kind of start that discourages questions. I "
        "want to ask a few anyway, because Arsenal's five matches do not tell one story. They "
        "tell two. In the first three, opponents barely got out of their own half. In the last "
        "two, both away, the same team with the same share of the ball was opened up "
        "repeatedly. The gap between those two versions is the most useful thing this early "
        "part of the season can teach us.\n"
        "Start with the outline of the side. Each slice of the wheel is one measure and its "
        "length is Arsenal's standing among the twenty clubs. The blue defensive slices are "
        "long, the amber build-up slices are healthy, and the red attacking ones disagree "
        "with each other.",
    ),
    (
        "overview_01",
        "The whole league on one page, ordered by points and shaded by rank in each column.",
        "Two things on that wheel look wrong at first glance, and they are the threads this "
        "piece follows. Arsenal get into the penalty area more than anyone, yet rank "
        "thirteenth for shots. And they post a top-five pressing number while sitting in the "
        "bottom six for winning the ball high. Neither is a contradiction once you see how "
        "the team is meant to work. Before that, it helps to see them next to everyone else.",
    ),
    (
        "club_match_by_match",
        "Chances created and conceded in each match, with the shot count and pressing intensity underneath.",
        "Read Arsenal's row left to right. The xGA and shots-against cells are among the "
        "brightest in their columns; the BOX cell is the brightest on the page; SH and HIGH "
        "are dull. Manchester City, three points clear, are brighter across the attacking "
        "columns and dimmer in defence. So this is a table position built from the back. But "
        "a row of averages hides the order things happened in, and the order matters.",
    ),
    (H, "Three games of control, two of something else"),
    (
        "overview_02",
        "xG created against xG conceded per match. Higher is better on both axes.",
        "Coventry, Aston Villa and Chelsea managed 1.26 expected goals between them. Villa "
        "did not have a shot on target. Then Sunderland produced 2.10 on their own and "
        "Brighton 1.70. Of the 5.06 xG Arsenal have conceded, 3.80 arrived in those two "
        "afternoons.\n"
        "The small print under the bars explains more than the bars do. PPDA counts the "
        "passes an opponent is allowed before a defensive action interrupts them, so lower "
        "means a harder press. Arsenal's reads 5.5, 5.3, 6.0, 7.3, 8.7. It has softened every "
        "week. As a block the first three games come out at 5.6, level with Liverpool as the "
        "third-hardest press in the league. The last two come out at 8.0, which would rank "
        "twelfth. Entries into Arsenal's box followed: 3, 7 and 7, then 12 and 11.\n"
        "The Sunderland result hid this. A 2-0 away win reads as routine, but Sunderland had "
        "the better chances and 0.79 of Arsenal's 1.92 xG was a penalty. Brighton removed the "
        "disguise. The season average will flatter one Arsenal and insult the other, but it "
        "still shows where they stand.",
    ),
    (H, "The defence: nothing gets through, and nobody is tackling"),
    (
        "defence_16",
        "Expected goals conceded per match, fewest at the top.",
        "Top right, with only Nottingham Forest for company at that height. Arsenal's xG "
        "difference of +0.69 a game is second to Manchester United's. Brighton create far "
        "more and pay for it at the other end; Forest concede as little and create less. The "
        "reason Arsenal sit where they do is almost entirely the vertical axis, so that is "
        "where to look first.",
    ),
    (
        "defence_01",
        "Goals conceded per match.",
        "1.01 expected goals against per game, a hair behind Forest and a long way inside "
        "the league median of 1.54. The model's view and the scoreboard's can differ over "
        "five games, so check the plain count.",
    ),
    (
        "defence_02",
        "Shots conceded per match.",
        "Four goals conceded, 0.8 a game, third-fewest. Three of the four came in one "
        "match. A low goals-against column can come from a good goalkeeper, a kind run of "
        "finishing, or from simply not letting the other side shoot. The next chart says "
        "which.",
    ),
    (
        "defence_03",
        "Shots on target conceded per match.",
        "10.6 shots against per game, second only to Forest, against a median near 14. "
        "Newcastle face nearly twice as many. And of the shots that are taken, few trouble "
        "the goal.",
    ),
    (
        "defence_04",
        "Post-shot expected goals conceded per match: the value of the shots that were on target.",
        "2.8 on target per game is the lowest figure in the league. Villa had none at all. "
        "Chelsea had five and scored once. When opponents shoot against Arsenal they tend to "
        "do it from distance or under pressure, which the post-shot model can put a number "
        "on.",
    ),
    (
        "defence_05",
        "Clear-cut chances conceded per match.",
        "0.73 xGOT against per game, second to Leeds. That is lower than the 1.01 xG "
        "those shots were worth before they were struck, so opponents have on balance made "
        "their chances worse with the finish. Some of that is pressure on the shooter and "
        "some is luck. The count of clear chances is harder to argue with.",
    ),
    (
        "defence_06",
        "Opposition entries into the penalty area per match.",
        "1.4 big chances against per game, second-fewest. But the split is brutal: one in "
        "the first three matches, six in the last two. Everything in the defensive record "
        "comes down to whether the opponent reaches the box in the first place.",
    ),
    (
        "defence_18",
        "Goals conceded against the xG of the chances given up.",
        "Eight box entries allowed per game, second only to United. That is the cause; the "
        "low shot count is the effect. Seventeen entries in the first three games, "
        "twenty-three in the last two. Now, did Arsenal get what they deserved from the "
        "chances they did give up?",
    ),
    (
        "defence_17",
        "Post-shot xG faced against goals conceded. Green means fewer goals than the shots deserved.",
        "Roughly, yes. Four conceded from 5.06 xG is one goal better than par, a small "
        "green bar next to Leeds and Everton, who are four ahead. This chart includes how "
        "badly opponents finished, though. To isolate the goalkeeper, start from the shots "
        "that were actually on target.",
    ),
    (
        "defence_07",
        "Goals prevented per match: post-shot xG faced minus goals conceded.",
        "Four goals from on-target shots worth 3.66. David Raya has conceded almost exactly "
        "what he faced. He has barely been asked a question in three of the five games, "
        "and in the other two he was par.",
    ),
    (
        "defence_19",
        "Tackles, interceptions, blocks and clearances per match for every club.",
        "Minus 0.07 a game, twelfth, dead on the league's middle. So the defensive record "
        "owes nothing to goalkeeping. Brighton and Everton have been bailed out; Sunderland "
        "have been let down. Arsenal's numbers are simply the chances they allowed. Which "
        "raises the obvious question of what the defenders are doing to allow so few.",
    ),
    (
        "defence_08",
        "Tackles per match.",
        "Very little, is the answer. The league's second-best defence by chances makes the "
        "second-fewest defensive actions: 43 a game, when Leeds and Crystal Palace need "
        "nearly 68. Take the four colours one at a time.",
    ),
    (
        "defence_09",
        "Interceptions per match.",
        "11.6 tackles per game, nineteenth. Palace make almost twice as many. A tackle is "
        "what happens when an attacker has got near enough to need stopping, and Arsenal's "
        "opponents mostly have not.",
    ),
    (
        "defence_21",
        "Tackles against interceptions per match.",
        "7.4 interceptions, fifteenth. You might expect a side that defends by position "
        "to read a lot of passes, but interceptions need the opponent to be passing in your "
        "half. Put the two together and Arsenal's place is clear.",
    ),
    (
        "defence_10",
        "Blocks per match.",
        "Bottom-left, in the corner the chart labels as little of either, alongside City "
        "and Brighton. It is the possession teams' corner. The sides in the top right, "
        "Palace, Leeds, Everton, are busy because they spend long spells without the ball.",
    ),
    (
        "defence_11",
        "Clearances per match.",
        "2.6 blocks per game; only Villa make fewer. Newcastle make 8.6. A block is a "
        "last-ditch act, a body in front of a shot. Arsenal's low number is the mirror "
        "of the low shot count above.",
    ),
    (
        "defence_14",
        "Aerial duels won, as a share of those contested.",
        "21.6 clearances, eighteenth. Newcastle clear the ball 35 times a game. The one "
        "match where Arsenal's defending was stretched, at Brighton, was actually their "
        "lowest clearance count of the season (15), because the chances came from play "
        "through them, not balls into the box. When the ball does go long, how do they do?",
    ),
    (
        "defence_15",
        "Ground duels won, as a share of those contested.",
        "50.7 per cent of aerial duels won, ninth. Average. For a side with this "
        "reputation at set pieces that is unremarkable, and it will come up again when we "
        "reach the corners. On the floor it is the same story.",
    ),
    (
        "defence_12",
        "Fouls committed per match.",
        "50.8 per cent of ground duels won, eighth. But the rate matters less than the "
        "volume: Arsenal contest 25 ground duels a game, the fewest in the league, where "
        "Tottenham contest 43. They are not winning the individual battles more often "
        "than anyone else. They are avoiding them. That shows in the discipline, too.",
    ),
    (
        "defence_13",
        "Yellow cards per match.",
        "10.2 fouls a game, fourteenth. Forest commit 15. A team that is rarely chasing "
        "has little reason to foul.",
    ),
    (
        "defence_20",
        "Fouls committed against yellow cards per match.",
        "1.6 bookings per game, in the lower half, and no red cards. The two measures "
        "together place them clearly.",
    ),
    (
        "defence_22",
        "Pressing intensity against chances conceded. Left is a harder press, down is a tighter defence.",
        "The clean corner, with City, Palace and Fulham nearby. So the picture of the "
        "defence is complete: few shots faced, few tackles made, few fouls needed, an "
        "ordinary goalkeeper's record. What protects Arsenal is where the game is played. "
        "The ball is the first line of defence, and the press that wins it back is the "
        "second. That makes the press a load-bearing wall.",
    ),
    (H, "The press: a fence, not a trap"),
    (
        "pressing_05",
        "PPDA for every club. Lower means fewer opponent passes allowed per defensive action.",
        "Pressing hard does not, by itself, keep you safe. Brighton and Tottenham press "
        "harder than anyone and concede around 1.7 a game, because a press that is beaten "
        "leaves a lot of grass behind it. Forest reach the same defensive number as Arsenal "
        "from the opposite direction, sitting off. Arsenal are alone in the bottom-left "
        "because they pair the press with possession: they are rarely asked to press for "
        "long.",
    ),
    (
        "pressing_08",
        "Pressing intensity against field tilt.",
        "Fifth in the league at 6.3, and that is the average of a top-three press and a "
        "bottom-half one. Where on the pitch it happens is the next thing to check.",
    ),
    (
        "pressing_01",
        "Possessions won in the final third per match.",
        "Top-left: hard press, high territory, in the cluster with City, Liverpool, "
        "Tottenham, United and Brighton. These are the sides that defend in the "
        "opposition half. You would expect every one of them to win the ball there often.",
    ),
    (
        "pressing_03",
        "Share of high regains that led to a shot.",
        "Arsenal do not. 14.6 high regains per game is in the bottom six, well short of "
        "Brighton's 21.8 and below the median of 18.4. This is the second oddity from the "
        "wheel: a top-five press that does not win the ball high. And what they do win "
        "there, they do little with.",
    ),
    (
        "pressing_07",
        "High regains per match, and the share of them that became a shot.",
        "5.3 per cent of Arsenal's high regains become a shot, seventeenth. Brighton turn "
        "nearly 12 per cent of theirs into one. Put volume and yield on the same chart.",
    ),
    (
        "pressing_02",
        "Possessions won back per match, anywhere on the pitch.",
        "Bottom-left, next to Villa. Brighton are alone in the far corner. Brighton's "
        "press is a trap, set to steal the ball near goal and attack a defence facing the "
        "wrong way. Arsenal's is a fence. Its job is to hurry the opponent into a long "
        "ball, collect the second phase, and start again with everyone in position. PPDA "
        "rewards that, because the opponent's passes are cut short. High regains do not, "
        "because the ball is won forty metres further back.",
    ),
    (
        "pressing_04",
        "Counter-press regains per match: the ball won back within seconds of losing it.",
        "68.8 regains a game, the lowest in the league. That sounds damning and is "
        "mostly arithmetic: Arsenal have the fewest regains partly because they give the "
        "ball away so rarely that there is less to win back. The same applies to the "
        "counter-press.",
    ),
    (
        "pressing_06",
        "Share of lost possessions won back within seconds.",
        "8.6 counter-press regains, sixteenth, from 74 attempts, nineteenth. Fewer losses, "
        "fewer chances to counter-press. The rate is the fairer measure.",
    ),
    (
        "pressing_09",
        "Counter-press success against total possessions won per match.",
        "11.7 per cent, thirteenth, below the league average. This one cannot be explained "
        "away by volume. When Arsenal do lose the ball, they win it straight back less "
        "often than City (14.8) or Brighton (17.1). For a team whose defence relies on "
        "staying compact around the ball, that is the softest number in the whole set.",
    ),
    (H, "The ball: round the block, rarely through it"),
    (
        "passing_11",
        "Share of possessions per match.",
        "Far left, low: the corner the chart calls passive. The label does not fit a side "
        "with the fifth-hardest press, but the position is real. Arsenal regain the ball "
        "less often than anyone and do not snap it back especially quickly. What they do "
        "is keep it. Everything in this section follows from that.",
    ),
    (
        "passing_01",
        "Passes attempted per match.",
        "59 per cent of possessions, fifth, behind Brighton, City, United and Tottenham. "
        "The spread across matches is wide: 67 per cent at Villa, 51 at home to Chelsea, "
        "which was also their best attacking display. More of the ball has not meant "
        "more chances.",
    ),
    (
        "passing_12",
        "Touches per match.",
        "540 passes a game, seventh. City play 666. Arsenal have a lot of the ball without "
        "being a high-tempo passing side: their games contain the fourth-fewest possessions "
        "in the league. They slow matches down.",
    ),
    (
        "passing_02",
        "Pass completion.",
        "604 touches, seventh, and 30 per cent of them in the attacking third, which is "
        "third-highest. The ball spends its time in the right half of the pitch.",
    ),
    (
        "passing_10",
        "Field tilt: the share of final-third touches in a team's matches that are its own.",
        "85.2 per cent completed, fourth. Safe, as you would expect. The more telling "
        "number is where those passes are played.",
    ),
    (
        "passing_19",
        "Pass completion against field tilt.",
        "64.8 per cent field tilt, fifth. Nearly two of every three final-third touches "
        "in an Arsenal match are Arsenal's. It was 72 per cent against Coventry and never "
        "below 61. Notably it did not drop in the two bad games. Arsenal had the territory "
        "at Sunderland and Brighton. They lost control of what happened when it was "
        "turned over.",
    ),
    (
        "passing_18",
        "Possession against xG difference per match.",
        "Top-right corner with City, Brighton, United, Tottenham and Liverpool: clean and "
        "high up. Six clubs play this way. They do not all get the same return from it.",
    ),
    (
        "passing_09",
        "Directness: share of ball progress made straight towards goal.",
        "Arsenal and United get the most from their possession, at around +0.7 xG a game. "
        "Tottenham have just as much of the ball and are half a goal a game in the red. "
        "Possession is a method, and Arsenal's version of it is working on aggregate. How "
        "they move the ball with it is the next question.",
    ),
    (
        "passing_16",
        "Pass completion against directness.",
        "33.6 per cent direct, fourteenth. Hull are at 60. Only a third of the ground "
        "Arsenal gain is gained straight up the pitch; the rest comes from moving the "
        "opponent sideways first.",
    ),
    (
        "passing_06",
        "Long balls per match.",
        "Bottom right, short and safe, with the same five clubs again. No surprise. The "
        "detail is in the exceptions to the short game.",
    ),
    (
        "passing_07",
        "Long-ball completion.",
        "44 long balls a game, eighteenth. They played 55 at Brighton, their most, "
        "because the league's most aggressive press gave them no choice.",
    ),
    (
        "passing_17",
        "Crosses against long balls per match.",
        "49 per cent of them completed, sixth. Arsenal go long rarely and accurately: a "
        "chosen pass, not a punt. That "
        "leaves the other way round a block.",
    ),
    (
        "passing_08",
        "Cross completion.",
        "Bottom-left: few crosses, few long balls. United and Tottenham keep the ball as "
        "much and cross far more. Arsenal's 14.8 crosses a game ranks sixteenth. For a "
        "side that reaches the byline so often, they rarely put it in the air.",
    ),
    (
        "passing_13",
        "Build-up success rate.",
        "23 per cent of crosses find a team-mate, eleventh. Average accuracy, low volume. "
        "The cross is not a route they trust, which fits a front line whose centre-forward "
        "wins aerials to lay the ball off, not to score. Back to the start of the move.",
    ),
    (
        "passing_03",
        "Progressive passes per match.",
        "37.6 per cent of build-ups succeed, fourth. Playing out from the back works. The "
        "goalkeeper has completed 17 line-breaking passes himself, second-most in the "
        "squad. Once out, how do they advance?",
    ),
    (
        "passing_04",
        "Completed line-breaking passes per match.",
        "Sixty progressive passes a game, fourth-highest, behind Liverpool, United and "
        "Brighton. On volume of forward passing Arsenal are among the best. On the "
        "sharper measure they are not.",
    ),
    (
        "passing_15",
        "Progressive passes on the left, completed line-breaking passes on the right, per match.",
        "22.4 completed line-breaking passes, exactly the league median. Liverpool "
        "complete 31. Set the two measures side by side.",
    ),
    (
        "passing_14",
        "Share of final-third possessions that went on to reach the penalty area.",
        "Arsenal's blue bar is long and their amber bar is ordinary. They advance the ball "
        "a great deal but mostly around the opposition's lines, not through them: out to "
        "the full-back, down the side, back inside. It is low-risk progression, in keeping "
        "with a team that would sooner keep the ball than gamble it between two "
        "midfielders. Their best line-breaking total of the season, 28, came in the "
        "defeat at Brighton, when the game was stretched and suited them least. The "
        "patient route does get them somewhere, though.",
    ),
    (
        "passing_05",
        "Key passes per match: the pass before a shot.",
        "First in the league. 26 per cent of Arsenal's final-third possessions end up in "
        "the penalty area, against a median of 17. Whatever happens before the final "
        "third, once there they are the most reliable side in the division at going the "
        "last step. The final pass is a different matter.",
    ),
    (H, "The attack: first to the box, thirteenth to the shot"),
    (
        "attack_11",
        "Entries into the final third per match.",
        "11.2 key passes a game, sixth. United play 16. So Arsenal's passing puts them in "
        "the box more than anyone but creates a shot only a little more often than "
        "average. Follow the ball up the pitch to see where the drop-off happens.",
    ),
    (
        "attack_10",
        "Entries into the penalty area per match.",
        "38.6 final-third entries, seventh. Liverpool and United are above 50. Arsenal do "
        "not arrive in the final third unusually often. They are unusually good at what "
        "comes next.",
    ),
    (
        "attack_12",
        "Deep completions per match.",
        "Sixteen box entries a game against a median of about ten. Nobody has more. They "
        "led this count when ahead, too, with six entries a game while leading: Arsenal "
        "do not sit on a lead.",
    ),
    (
        "attack_13",
        "Expected threat added per match.",
        "13 deep completions, second to United. The ball is being received near goal "
        "regularly, which the threat model also picks up.",
    ),
    (
        "attack_22",
        "Key passes against passes into the box per match.",
        "7.08 xT a game, third. Expected threat values where the ball is moved to, not "
        "whether a shot follows. On location, Arsenal's possession is elite. The gap "
        "opens between location and shot.",
    ),
    (
        "attack_21",
        "Box entries per match against expected goals per match.",
        "Furthest right, only just above the middle. United are the comparison: similar "
        "passes into the box, five more key passes a game. Arsenal's entries are often "
        "passes to a player facing away from goal or pinned on the byline, with one more "
        "pass needed before anyone can shoot.",
    ),
    (
        "attack_14",
        "Crosses per match.",
        "Furthest right on this chart too, and well short of the top. United create 2.12 "
        "xG from almost the same entries; Brighton create 2.35 from about eleven. Just "
        "under half of Arsenal's box entries lead to a shot, fifteenth. Their entries "
        "arrive wide and against a defence that has already retreated, because the slow "
        "build-up gave it time to. Two routes could turn that territory into shots more "
        "cheaply. Arsenal use neither much.",
    ),
    (
        "attack_15",
        "Corners won per match.",
        "Sixteenth for crosses. The ball goes to the byline and comes back along the floor.",
    ),
    (
        "attack_04",
        "Shots per match.",
        "8.4 corners, fifteenth. All that territory forces fewer corners than Tottenham or "
        "Chelsea win, because Arsenal's final action is so often a short pass that gets "
        "recycled, not a shot or cross that gets deflected behind. So the shot count is "
        "what you would predict.",
    ),
    (
        "attack_08",
        "Average xG per shot.",
        "Thirteen a game, thirteenth. United take 22.6. The range inside Arsenal's own "
        "season is striking: 20 against Coventry, 7 at Villa. Those 7 were enough to win. "
        "The reason is the quality.",
    ),
    (
        "attack_20",
        "Shots per match against the average xG of each shot.",
        "0.131 xG per shot, third behind Sunderland and City. Arsenal turn down the "
        "speculative effort.",
    ),
    (
        "attack_05",
        "Shots on target per match.",
        "The patient corner: below-average volume, high quality. Part of the low "
        "entry-to-shot rate is therefore a choice, and mostly a sound one. The risk is "
        "that it is a habit as well. Chasing the game at Brighton, they finished with "
        "eleven shots to the home side's seventeen. Better shots should mean more of "
        "them hit the target.",
    ),
    (
        "attack_06",
        "Shot accuracy: share of shots on target.",
        "4.8 on target per game, seventh, from the thirteenth-most shots.",
    ),
    (
        "attack_02",
        "Expected goals per match.",
        "37 per cent on target, fifth. They are also third-lowest for shots off target. "
        "Close-range shots are hard to miss the frame with. Add volume and quality "
        "together.",
    ),
    (
        "attack_03",
        "Post-shot expected goals per match: the value of the shots after they were struck.",
        "1.70 xG a game, sixth, behind Brighton, United, Sunderland, City and Brentford. "
        "Good, not outstanding, for a side with this much of the ball. What happens to "
        "those chances once they are struck?",
    ),
    (
        "attack_09",
        "Clear-cut chances created per match.",
        "1.47, eighth, and lower than the 1.70 the chances were worth beforehand. "
        "Arsenal's finishing has taken value off their shots: well-placed positions, "
        "ordinary placement. Brighton's has added to theirs. It is a small-sample "
        "number, but it argues against the idea that Arsenal have been unlucky in "
        "front of goal.",
    ),
    (
        "attack_18",
        "Clear-cut chances created and conceded per match, ordered by the balance of the two.",
        "2.4 big chances a game, ninth. Brighton create 4.4. On the provider's own "
        "definition of a clear chance, Arsenal are mid-table creators. Set against what "
        "they concede, the picture changes.",
    ),
    (
        "attack_19",
        "Expected goals per match, split by how the move began.",
        "Top of the chart. 2.4 created to 1.4 conceded is the best balance in the league, "
        "with Forest next. Arsenal do not win this by creating the most. They win it by "
        "conceding the fewest, which is the whole season in one bar. Where the 1.70 comes "
        "from adds a qualification.",
    ),
    (
        "attack_16",
        "Expected goals from transitions per match.",
        "Look at the blue section alone and Arsenal drop to tenth: 1.17 xG from open "
        "play, a whisker above the median. The red section lifts them, and it is thinner "
        "than it looks. Free-kick situations against Chelsea were worth 0.8 xG and the "
        "penalty at Sunderland 0.79; those two matches' dead balls are nearly a fifth of "
        "everything Arsenal have created. Corners, the part of their game with the "
        "reputation, are tenth for xG with no goals. What the territory does reliably "
        "produce is fouls: Arsenal win over fifteen free kicks a game, the most in the "
        "league. One source is missing from that chart altogether.",
    ),
    (
        "attack_01",
        "Goals scored per match.",
        "0.44 xG a game from transitions, fifteenth, from the second-fewest transition "
        "shots. Brentford get 1.11. Arsenal have nearly designed the counter-attack out "
        "of their game: they rarely win the ball in a position to break, and when they "
        "do they tend to keep it. Giving up the cheap chance a turnover offers is the "
        "cost of the fence. So to the output.",
    ),
    (
        "attack_07",
        "Goal conversion: goals as a share of shots.",
        "Eight goals, 1.6 a game, sixth. Brighton have twice as many. Three of the eight "
        "came against Coventry.",
    ),
    (
        "attack_17",
        "Goals scored against expected goals for every club.",
        "12.3 per cent of shots scored, sixth. A high conversion rate normally warns of a "
        "hot streak. Here it reflects shot selection: the shots are good ones.",
    ),
    (H, "What to watch from here"),
    (
        None,
        None,
        "Eight goals from 8.5 xG is as close to par as it gets. No hot streak, no "
        "goalkeeper standing on his head at the other end. Brighton are four goals ahead "
        "of their chances and Sunderland four behind. Arsenal's twelve points were earned "
        "by the performances, which is reassuring about the first three games and gives "
        "no excuse for the last two.\n"
        "Put the eighty charts together and they describe a chain. Keep the ball, hold "
        "territory, press quickly enough to get it back, and the back line is left with "
        "very little to do. Each link depends on the one before. When the press let "
        "Sunderland and Brighton string passes together, there was no second layer of "
        "duels and tackles underneath, because the team is not built to need one. Only "
        "two sides have faced more dangerous counter-attacks per game.\n"
        "The same chain explains the attack. Slow, safe progression gets Arsenal into the "
        "box more than anyone, and gives the defence time to be set when they arrive. "
        "They wait for a good shot, so they take few. They do not cross, do not counter, "
        "and have not yet made their corners count. What is left is an open-play attack "
        "of average output, lifted by a penalty and a set of free kicks.\n"
        "Three things will tell us which Arsenal this is. Whether the press travels: 5.6 "
        "in the first three games, 8.0 in the two away games that followed. Whether the "
        "counter-press, the softest number they have, tightens up. And whether open-play "
        "creation moves off the league average once the big dead-ball moments are taken "
        "out. Five matches is a lean, not a verdict. But the lean is clear enough: this "
        "is a very good side with one way of playing, and the last fortnight showed what "
        "the other way looks like.",
    ),
]

METHOD = (
    "About the numbers. Everything here covers Premier League rounds 1 to 5 of 2026-27, fifty "
    "matches, from the WhoScored event feed. Expected goals come from my own shot model, so "
    "they will differ a little from other providers. Team figures are per match; with five "
    "matches played, every one of them is a small sample."
)


def picture_for(folder: Path, prefix: str) -> Path:
    """The one image in ``folder`` whose name starts with ``prefix``."""
    found = sorted(folder.glob(f"{prefix}*.png"))
    if len(found) != 1:
        raise FileNotFoundError(f"{prefix}: {len(found)} images in {folder}")
    return found[0]


def build(folder: Path = FOLDER) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Inches(1)
    section.top_margin = section.bottom_margin = Inches(1)
    width = section.page_width - section.left_margin - section.right_margin

    body = doc.styles["Normal"]
    body.font.name = "Georgia"
    body.font.size = Pt(11.5)
    body.paragraph_format.space_after = Pt(10)
    body.paragraph_format.line_spacing = 1.35
    for name, size in (("Title", 26), ("Heading 2", 15)):
        style = doc.styles[name]
        style.font.name = "Georgia"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(20)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(8)

    doc.core_properties.title = TITLE
    doc.core_properties.author = "Mostafa Saad"
    doc.add_paragraph(TITLE, style="Title").alignment = WD_ALIGN_PARAGRAPH.LEFT
    stand = doc.add_paragraph()
    run = stand.add_run(STANDFIRST)
    run.italic = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    stand.paragraph_format.space_after = Pt(16)

    for block in BLOCKS:
        if block[0] == H:
            doc.add_paragraph(block[1], style="Heading 2")
            continue
        image, caption, text = block
        for paragraph in text.split("\n"):
            doc.add_paragraph(paragraph)
        if image is None:
            continue
        figure = doc.add_paragraph()
        figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
        figure.paragraph_format.space_after = Pt(4)
        figure.paragraph_format.keep_with_next = True
        figure.add_run().add_picture(str(picture_for(folder, image)), width=width)
        note = doc.add_paragraph()
        note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        note.paragraph_format.space_after = Pt(14)
        words = note.add_run(caption)
        words.italic = True
        words.font.size = Pt(9.5)
        words.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    method = doc.add_paragraph()
    method.paragraph_format.space_before = Pt(18)
    small = method.add_run(METHOD)
    small.font.size = Pt(9.5)
    small.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    target = folder / "Arsenal_tactical_analysis_R1-5.docx"
    doc.save(target)
    return target


if __name__ == "__main__":
    path = build()
    used = [b[0] for b in BLOCKS if b[0] not in (H, None)]
    unused = sorted(
        p.stem for p in FOLDER.glob("*.png") if not any(p.name.startswith(u) for u in used)
    )
    words = sum(len(b[2].split()) for b in BLOCKS if b[0] != H)
    print(path, f"{words} words, {len(used)} images, {len(set(used))} distinct, unused: {unused}")
